"""
Parse a TrackEval-style results .txt and append the metrics to an Excel sheet.

Columns written:
    Detector | Tracker | HOTA | DetA | DetRe | DetPr | AssA | AssRe | AssPr | IDF1 |
    Det | GT_Det | ID | GT_ID | |ID-GT_ID| | IDSW | Notes

Usage (PowerShell):
    python .\results_to_excel.py .\trackeval_output_for_normal_yolov8n.txt
    python .\results_to_excel.py .\a.txt -d "YOLOv8n" -o sheep_results.xlsx
    python .\results_to_excel.py .\a.txt .\b.txt -o sheep_results.xlsx
    
    eg: python .\results_to_excel.py .\trackeval_output_for_normal_yolov8n.txt -d "YOLOv8n" -o rfdetr_sheep_results.xlsx

    Process EVERY .txt in the trackeval_outputs folder (detector name taken from each file name):
    python .\results_to_excel.py
    python .\results_to_excel.py --folder D:\\other\\trackeval_outputs -o sheep_results.xlsx

Running it again with another txt APPENDS rows to the same workbook, so you can
build up one table across detectors. The Notes column is left empty for you.
"""
import argparse
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Same convention as the other scripts: trackeval_outputs sits in the folder above this script
OUTPUTS_DIR = Path(__file__).resolve().parent.parent / "trackeval_outputs"

HEADERS = ["Detector", "Tracker",
           "HOTA", "DetA", "DetRe", "DetPr", "AssA", "AssRe", "AssPr", "IDF1",
           "Det", "GT_Det", "ID", "GT_ID", "|ID-GT_ID|", "IDSW", "Notes"]

# Columns that come straight from the txt (|ID-GT_ID| is a formula, Notes is blank)
DATA_COLS = ["HOTA", "DetA", "DetRe", "DetPr", "AssA", "AssRe", "AssPr", "IDF1",
             "Det", "GT_Det", "ID", "GT_ID"]
FLOAT_COLS = {"HOTA", "DetA", "DetRe", "DetPr", "AssA", "AssRe", "AssPr", "IDF1"}
COUNT_COLS = {"Det", "GT_Det", "ID", "GT_ID", "|ID-GT_ID|", "IDSW"}

# (section in txt) -> {column name in txt: column name in Excel}
WANTED = {
    "HOTA":     {c: c for c in ("HOTA", "DetA", "DetRe", "DetPr", "AssA", "AssRe", "AssPr")},
    "Identity": {"IDF1": "IDF1"},
    "Count":    {"Dets": "Det", "GT_Dets": "GT_Det", "IDs": "ID", "GT_IDs": "GT_ID"},
    "CLEAR":    {"IDSW": "IDSW"},
}

CLASS_SUFFIX = re.compile(r"-(pedestrian|car|person|vehicle|all)$", re.I)
SECTION_RE = re.compile(r"^(HOTA|CLEAR|Identity|Count):\s+(\S+)\s+(.*)$")


def col_letter(name):
    return get_column_letter(HEADERS.index(name) + 1)


def parse_txt(path):
    """Return {tracker_name: {excel_col: value}} for every tracker in the file."""
    lines = Path(path).read_text(errors="ignore").splitlines()
    results = {}
    for i, line in enumerate(lines):
        m = SECTION_RE.match(line.strip())
        if not m:
            continue
        section, tracker, cols = m.group(1), m.group(2), m.group(3).split()
        tracker = CLASS_SUFFIX.sub("", tracker)
        # find the COMBINED row that belongs to this header
        for nxt in lines[i + 1:]:
            parts = nxt.split()
            if parts and parts[0] == "COMBINED":
                vals = parts[1:]
                for src, dst in WANTED[section].items():
                    if src in cols and cols.index(src) < len(vals):
                        results.setdefault(tracker, {})[dst] = float(vals[cols.index(src)])
                break
            if SECTION_RE.match(nxt.strip()):
                break  # hit next section without finding COMBINED
    return results


def open_workbook(path):
    if Path(path).exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        ws.append(HEADERS)
        hfont = Font(name="Arial", bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="305496")
        for c in ws[1]:
            c.font, c.fill = hfont, fill
            c.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
    return wb, ws


def style_and_size(ws):
    thin = Side(style="thin", color="BFBFBF")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            name = HEADERS[c.column - 1]
            c.font = Font(name="Arial", size=10)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if name in FLOAT_COLS:
                c.number_format = "0.000"
            elif name in COUNT_COLS:
                c.number_format = "#,##0"
            elif name == "Notes":
                c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {"Detector": 20, "Tracker": 26, "|ID-GT_ID|": 14, "Notes": 40}
    for idx, name in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 11)


def detector_name_from(txt):
    return re.sub(r"^trackeval_output_for_", "", Path(txt).stem)


def collect_txt_files(folder=OUTPUTS_DIR):
    """Every .txt in the trackeval_outputs folder, sorted by name."""
    folder = Path(folder)
    if not folder.is_dir():
        raise FileNotFoundError(f"Folder not found: {folder}")
    files = sorted(folder.glob("*.txt"))
    if not files:
        raise FileNotFoundError(f"No .txt files found in: {folder}")
    return files


def add_results(ws, txt_files, detector=None, skip_existing=False):
    """Parse each txt and append its rows to the sheet. Returns number of rows added."""
    id_c, gt_c = col_letter("ID"), col_letter("GT_ID")
    # (detector, tracker) pairs already in the sheet, to avoid duplicates when re-running
    existing = {(r[0], r[1]) for r in ws.iter_rows(min_row=2, max_col=2, values_only=True)}
    added = 0

    for txt in txt_files:
        name = detector or detector_name_from(txt)
        try:
            data = parse_txt(txt)
        except OSError as e:
            print(f"[!] Could not read {txt}: {e}")
            continue
        if not data:
            print(f"[!] No metrics found in {txt} - check the file format.")
            continue

        n = 0
        for tracker, m in data.items():
            if skip_existing and (name, tracker) in existing:
                continue
            row = ws.max_row + 1
            ws.append([name, tracker]
                      + [m.get(h) for h in DATA_COLS]
                      + [f"=ABS({id_c}{row}-{gt_c}{row})", m.get("IDSW"), ""])
            existing.add((name, tracker))
            n += 1

        if n:
            print(f"[+] {Path(txt).name}: {n} tracker(s) added under detector '{name}'")
        else:
            print(f"[=] {Path(txt).name}: already in the sheet, skipped")
        added += n
    return added

import os
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("txt", nargs="*", help="results .txt file(s); if none are given, "
                                           "every .txt in the trackeval_outputs folder is used")
    ap.add_argument("-o", "--output", default=os.path.join(os.path.dirname(os.path.abspath(OUTPUTS_DIR)), "sheep_results.xlsx"))
    ap.add_argument("-d", "--detector", help="detector name (only valid with a single txt; "
                                             "default = txt file name)")
    ap.add_argument("-f", "--folder", default=str(OUTPUTS_DIR),
                    help="folder used when no txt files are given "
                         f"(default: {OUTPUTS_DIR})")
    args = ap.parse_args()
    if args.detector and len(args.txt) > 1:
        ap.error("--detector can only be used with one txt file")

    folder_mode = not args.txt
    if folder_mode:
        if args.detector:
            ap.error("--detector can't be used when processing a whole folder")
        try:
            txt_files = collect_txt_files(args.folder)
        except FileNotFoundError as e:
            ap.error(str(e))
        print(f"Processing {len(txt_files)} file(s) from {args.folder}")
    else:
        txt_files = args.txt

    try:
        wb, ws = open_workbook(args.output)
    except Exception as e:
        ap.error(f"Could not open {args.output}: {e} (is it open in Excel?)")

    # In folder mode, skip detector/tracker rows already in the workbook so re-runs don't duplicate
    add_results(ws, txt_files, detector=args.detector, skip_existing=folder_mode)

    style_and_size(ws)
    try:
        print(args.output)
        wb.save(args.output)
    except PermissionError:
        ap.error(f"Cannot write {args.output} - close it in Excel and try again.")
    print(f"Saved -> {args.output}")


if __name__ == "__main__":
    main()